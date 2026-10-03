import openpyxl
import os

FILE_NAME = "banking_system.xlsx"


# Find Account
def find_account(account_no):
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value == account_no:
            return workbook, sheet, row

    workbook.close()
    return None, None, None


# Verify PIN
def verify_pin(sheet, row):
    pin = input("Enter 4-digit PIN: ")

    if str(sheet.cell(row, 4).value) == pin:
        return True
    else:
        print("Incorrect PIN!")
        return False


# Create Account
def create_account():
    workbook = openpyxl.load_workbook(FILE_NAME)
    sheet = workbook.active

    name = input("Enter account holder name: ")
    mobile = input("Enter mobile number: ")
    pin = input("Enter 4-digit PIN: ")
    initial_deposit = float(input("Enter initial deposit: "))

    if len(pin) != 4 or not pin.isdigit():
        print("PIN must be exactly 4 digits!")
        workbook.close()
        return

    if initial_deposit < 0:
        print("Invalid initial deposit!")
        workbook.close()
        return

    # Generate unique account number
    account_no = 1001

    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value is not None:
            account_no = max(
                account_no,
                int(sheet.cell(row, 1).value) + 1
            )

    sheet.append([
        account_no,
        name,
        mobile,
        pin,
        initial_deposit
    ])

    workbook.save(FILE_NAME)
    workbook.close()

    print("Account created successfully!")
    print("Your Account Number is:", account_no)


# Deposit Money
def deposit_money():
    account_no = int(input("Enter account number: "))

    workbook, sheet, row = find_account(account_no)

    if row is None:
        print("Account not found!")
        return

    if not verify_pin(sheet, row):
        workbook.close()
        return

    amount = float(input("Enter deposit amount: "))

    if amount <= 0:
        print("Invalid amount!")
        workbook.close()
        return

    balance = sheet.cell(row, 5).value
    balance = balance + amount

    sheet.cell(row, 5).value = balance

    workbook.save(FILE_NAME)
    workbook.close()

    print("Money deposited successfully!")
    print("Updated Balance:", balance)


# Withdraw Money
def withdraw_money():
    account_no = int(input("Enter account number: "))

    workbook, sheet, row = find_account(account_no)

    if row is None:
        print("Account not found!")
        return

    if not verify_pin(sheet, row):
        workbook.close()
        return

    amount = float(input("Enter withdrawal amount: "))

    if amount <= 0:
        print("Invalid amount!")
        workbook.close()
        return

    balance = sheet.cell(row, 5).value

    if amount > balance:
        print("Insufficient balance!")
        workbook.close()
        return

    balance = balance - amount
    sheet.cell(row, 5).value = balance

    workbook.save(FILE_NAME)
    workbook.close()

    print("Money withdrawn successfully!")
    print("Updated Balance:", balance)


# Check Balance
def check_balance():
    account_no = int(input("Enter account number: "))

    workbook, sheet, row = find_account(account_no)

    if row is None:
        print("Account not found!")
        return

    if not verify_pin(sheet, row):
        workbook.close()
        return

    print("\n----- ACCOUNT DETAILS -----")
    print("Account Number :", sheet.cell(row, 1).value)
    print("Account Holder :", sheet.cell(row, 2).value)
    print("Current Balance:", sheet.cell(row, 5).value)

    workbook.close()


# Main Function
def main():

    # Create empty Excel file if it does not exist
    if not os.path.exists(FILE_NAME):

        workbook = openpyxl.Workbook()
        sheet = workbook.active

        sheet.append([
            "Account Number",
            "Name",
            "Mobile Number",
            "PIN",
            "Balance"
        ])

        workbook.save(FILE_NAME)
        workbook.close()

    # Main Menu
    while True:

        print("\n========== BANKING SYSTEM ==========")
        print("1. Create Account")
        print("2. Deposit Money")
        print("3. Withdraw Money")
        print("4. Check Balance")
        print("5. Exit")

        choice = input("Enter your choice: ")

        if choice == "1":
            create_account()

        elif choice == "2":
            deposit_money()

        elif choice == "3":
            withdraw_money()

        elif choice == "4":
            check_balance()

        elif choice == "5":
            print("Thank you for using Banking System!")
            break

        else:
            print("Invalid choice! Please try again.")


main()
